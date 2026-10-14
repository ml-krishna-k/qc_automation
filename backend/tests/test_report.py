"""Regression tests for the report/UI data contract: added/removed rows must
be labeled clearly ("Added Tag"/"Removed Tag") with no repeated tag-ID value,
matching the frontend's results table and its report builder
(`frontend/src/utils/differenceDisplay.ts`). Changed/error rows are
untouched — only presentation of added/removed rows changes, not the
comparison engine's output."""

from io import BytesIO

from openpyxl import load_workbook

from app.models import ComparisonConfig, ComparisonResult, ComparisonSummary, DifferenceRow
from app.report import build_comparison_report


def _make_result() -> ComparisonResult:
    return ComparisonResult(
        id="run_1",
        runAt="2026-01-01T00:00:00+00:00",
        config=ComparisonConfig(originalSheet="Original", revisedSheet="Revised", tagColumn="Tag Number"),
        summary=ComparisonSummary(
            totalTagsCompared=3,
            changedCount=1,
            addedCount=1,
            removedCount=1,
            errorCount=0,
            unchangedCount=0,
            columnsCompared=1,
        ),
        differences=[
            DifferenceRow(
                id="diff_changed",
                tagNumber="PV-1001",
                columnName="Status",
                originalValue="Approved",
                revisedValue="Superseded",
                changeType="changed",
            ),
            DifferenceRow(
                id="diff_added",
                tagNumber="PV-1003",
                columnName="Tag Number",
                originalValue=None,
                revisedValue="PV-1003",
                changeType="added",
            ),
            DifferenceRow(
                id="diff_removed",
                tagNumber="PV-1002",
                columnName="Tag Number",
                originalValue="PV-1002",
                revisedValue=None,
                changeType="removed",
            ),
        ],
    )


def _row_by_tag(sheet, tag_number: str) -> dict:
    headers = [cell.value for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        if record["Tag Number"] == tag_number:
            return record
    raise AssertionError(f"No row found for tag {tag_number}")


def test_added_and_removed_rows_are_labeled_clearly_without_repeating_the_tag_id():
    report_bytes = build_comparison_report(_make_result())
    workbook = load_workbook(BytesIO(report_bytes))
    diff_sheet = workbook["Differences"]

    changed = _row_by_tag(diff_sheet, "PV-1001")
    assert changed["Column"] == "Status"
    assert changed["Original Value"] == "Approved"
    assert changed["Revised Value"] == "Superseded"

    added = _row_by_tag(diff_sheet, "PV-1003")
    assert added["Column"] == "Added Tag"
    # openpyxl round-trips a written "" cell as either "" or None — both mean "blank".
    assert added["Original Value"] in ("", None)
    assert added["Revised Value"] in ("", None)

    removed = _row_by_tag(diff_sheet, "PV-1002")
    assert removed["Column"] == "Removed Tag"
    assert removed["Original Value"] in ("", None)
    assert removed["Revised Value"] in ("", None)
