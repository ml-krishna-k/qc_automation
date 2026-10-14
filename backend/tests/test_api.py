from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import config
from app.main import app
from tests.helpers import build_xlsx_bytes

client = TestClient(app)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client: TestClient, data: bytes, filename: str, slot: str) -> dict:
    response = client.post(
        "/api/workbooks",
        files={"file": (filename, data, XLSX_CONTENT_TYPE)},
        data={"slot": slot},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_full_comparison_flow():
    original_bytes = build_xlsx_bytes(
        ["Tag Number", "Description", "Status"],
        [
            ["PV-1001", "Pressure Vessel", "Approved"],
            ["PV-1002", "Separator", "Approved"],
        ],
    )
    revised_bytes = build_xlsx_bytes(
        ["Tag Number", "Description", "Status"],
        [
            ["PV-1001", "Pressure Vessel", "Superseded"],  # changed
            ["PV-1003", "Surge Drum", "Approved"],  # added (PV-1002 removed)
        ],
    )

    original = _upload(client, original_bytes, "original.xlsx", "original")
    revised = _upload(client, revised_bytes, "revised.xlsx", "revised")

    assert original["sheetNames"] == ["Sheet1"]

    columns_response = client.get(f"/api/workbooks/{original['id']}/columns", params={"sheet": "Sheet1"})
    assert columns_response.status_code == 200
    assert columns_response.json() == ["Tag Number", "Description", "Status"]

    inspect_response = client.get(f"/api/workbooks/{original['id']}/inspect", params={"sheet": "Sheet1"})
    assert inspect_response.status_code == 200
    assert inspect_response.json()["detectedTagColumn"] == "Tag Number"

    comparison_payload = {
        "originalWorkbookId": original["id"],
        "revisedWorkbookId": revised["id"],
        "config": {
            "originalSheet": "Sheet1",
            "revisedSheet": "Sheet1",
            "tagColumn": "Tag Number",
            "ignoredColumns": [],
            "trimWhitespace": True,
            "ignoreCase": False,
        },
    }
    comparison_response = client.post("/api/comparisons", json=comparison_payload)
    assert comparison_response.status_code == 200, comparison_response.text
    result = comparison_response.json()

    assert result["summary"]["changedCount"] == 1
    assert result["summary"]["addedCount"] == 1
    assert result["summary"]["removedCount"] == 1

    report_response = client.get(f"/api/comparisons/{result['id']}/report")
    assert report_response.status_code == 200
    assert report_response.headers["content-type"] == XLSX_CONTENT_TYPE

    report_workbook = load_workbook(BytesIO(report_response.content))
    assert report_workbook.sheetnames == ["Summary", "Differences"]
    assert report_workbook["Differences"].max_row == 1 + len(result["differences"])


def test_upload_rejects_unsupported_extension():
    response = client.post(
        "/api/workbooks",
        files={"file": ("notes.txt", b"hello", "text/plain")},
        data={"slot": "original"},
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["message"]


def test_upload_rejects_oversized_file(monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_SIZE_BYTES", 10)
    data = build_xlsx_bytes(["Tag Number"], [["T-1"]])
    response = client.post(
        "/api/workbooks",
        files={"file": ("original.xlsx", data, XLSX_CONTENT_TYPE)},
        data={"slot": "original"},
    )
    assert response.status_code == 413
    assert "too large" in response.json()["message"]


def test_unknown_workbook_id_returns_404():
    response = client.get("/api/workbooks/does-not-exist/columns", params={"sheet": "Sheet1"})
    assert response.status_code == 404
    assert "no longer available" in response.json()["message"]


def test_unknown_sheet_name_returns_404():
    data = build_xlsx_bytes(["Tag Number"], [["T-1"]])
    uploaded = _upload(client, data, "original.xlsx", "original")
    response = client.get(f"/api/workbooks/{uploaded['id']}/columns", params={"sheet": "Nope"})
    assert response.status_code == 404


def test_ambiguous_tag_column_returns_422_with_candidates():
    data = build_xlsx_bytes(["Asset Tag", "Equipment Tag"], [["T-1", "x"]])
    original = _upload(client, data, "original.xlsx", "original")
    revised = _upload(client, data, "revised.xlsx", "revised")

    response = client.post(
        "/api/comparisons",
        json={
            "originalWorkbookId": original["id"],
            "revisedWorkbookId": revised["id"],
            "config": {
                "originalSheet": "Sheet1",
                "revisedSheet": "Sheet1",
                "tagColumn": "Tag Number",
                "ignoredColumns": [],
                "trimWhitespace": True,
                "ignoreCase": False,
            },
        },
    )
    assert response.status_code == 422
    body = response.json()
    assert set(body["candidates"]) == {"Asset Tag", "Equipment Tag"}


def test_comparison_with_missing_workbook_returns_404():
    response = client.post(
        "/api/comparisons",
        json={
            "originalWorkbookId": "missing",
            "revisedWorkbookId": "also-missing",
            "config": {
                "originalSheet": "Sheet1",
                "revisedSheet": "Sheet1",
                "tagColumn": "Tag Number",
                "ignoredColumns": [],
                "trimWhitespace": True,
                "ignoreCase": False,
            },
        },
    )
    assert response.status_code == 404
