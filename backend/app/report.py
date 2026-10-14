"""Builds the downloadable .xlsx report (Summary + Differences sheets) from
a ComparisonResult. Mirrors the shape of the frontend's local report writer
(`frontend/src/api/buildReport.ts`) so a user sees the same layout regardless
of which engine produced the result."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.models import ComparisonResult, DifferenceRow

_HEADER_FONT = Font(bold=True)


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT


def _display_difference(diff: DifferenceRow) -> tuple[str, str, str]:
    """Presentation-only mapping, mirroring the frontend's
    `toDisplayedDifference` (`frontend/src/utils/differenceDisplay.ts`), so
    the downloaded report matches what the UI shows.

    The comparison engine represents an added/removed tag by reusing the tag
    column (e.g. "Tag Number") with the tag ID as the value — accurate, but
    next to the row's own Tag Number cell it reads as a duplicated, confusing
    "value change". A `changed`/`error` row's column and values are already
    meaningful and are left as-is.
    """
    if diff.changeType.value == "added":
        return "Added Tag", "", ""
    if diff.changeType.value == "removed":
        return "Removed Tag", "", ""
    return diff.columnName, diff.originalValue or "", diff.revisedValue or ""


def build_comparison_report(result: ComparisonResult) -> bytes:
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _write_header(summary_sheet, ["Metric", "Value"])
    summary_rows = [
        ("Run date", result.runAt),
        ("Original sheet", result.config.originalSheet),
        ("Revised sheet", result.config.revisedSheet),
        ("Tag column", result.config.tagColumn),
        ("Tags compared", result.summary.totalTagsCompared),
        ("Columns compared", result.summary.columnsCompared),
        ("Changed values", result.summary.changedCount),
        ("Added tags", result.summary.addedCount),
        ("Removed tags", result.summary.removedCount),
        ("Validation errors", result.summary.errorCount),
        ("Unchanged tags", result.summary.unchangedCount),
    ]
    for row in summary_rows:
        summary_sheet.append(row)
    summary_sheet.column_dimensions["A"].width = 20
    summary_sheet.column_dimensions["B"].width = 40

    diff_sheet = workbook.create_sheet("Differences")
    _write_header(
        diff_sheet,
        ["Tag Number", "Column", "Original Value", "Revised Value", "Change Type", "Original Row", "Revised Row", "Notes"],
    )
    for diff in result.differences:
        column_label, original_display, revised_display = _display_difference(diff)
        diff_sheet.append(
            [
                diff.tagNumber,
                column_label,
                original_display,
                revised_display,
                diff.changeType.value,
                diff.originalRow or "",
                diff.revisedRow or "",
                diff.errorMessage or "",
            ]
        )
    widths = [14, 20, 28, 28, 12, 12, 12, 48]
    for idx, width in enumerate(widths, start=1):
        diff_sheet.column_dimensions[diff_sheet.cell(row=1, column=idx).column_letter].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
