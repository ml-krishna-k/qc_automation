"""openpyxl-backed reading of uploaded workbooks. Cell values are kept in
their native Python type (str/int/float/bool/datetime/None) so comparisons
elsewhere are exact — no stringification happens until display time."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


class WorkbookReadError(ValueError):
    """Raised when a file can't be parsed as an Excel workbook."""


class SheetNotFoundError(ValueError):
    """Raised when a requested worksheet name doesn't exist in the workbook."""


def load_workbook_from_path(path: Path) -> Workbook:
    try:
        return load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # openpyxl raises several distinct exception types
        raise WorkbookReadError(f"Could not read this file as an Excel workbook: {exc}") from exc


def load_workbook_from_bytes(data: bytes) -> Workbook:
    try:
        return load_workbook(BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:
        raise WorkbookReadError(f"Could not read this file as an Excel workbook: {exc}") from exc


@dataclass
class SheetRow:
    row_number: int
    values: dict[str, Any]


@dataclass
class ParsedSheet:
    columns: list[str]  # unique, in file order
    duplicate_headers: list[str]
    rows: list[SheetRow]


def get_sheet(workbook: Workbook, sheet_name: str):
    if sheet_name not in workbook.sheetnames:
        raise SheetNotFoundError(f'Worksheet "{sheet_name}" was not found in this workbook.')
    return workbook[sheet_name]


def read_sheet_columns(workbook: Workbook, sheet_name: str) -> list[str]:
    """Unique, ordered header names from row 1 (blank header cells skipped)."""
    sheet = get_sheet(workbook, sheet_name)
    header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
    seen: dict[str, None] = {}
    for cell in header_row:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if name:
            seen.setdefault(name, None)
    return list(seen.keys())


def parse_sheet(workbook: Workbook, sheet_name: str) -> ParsedSheet:
    """Reads a worksheet into row dicts keyed by header name, tracking:
    - duplicate header names (a data-quality issue reported to the user)
    - the original spreadsheet row number of every data row
    - fully blank rows, which are skipped entirely
    """
    sheet = get_sheet(workbook, sheet_name)

    header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
    column_order: list[str] = []
    seen_headers: set[str] = set()
    duplicate_headers: set[str] = set()
    col_index_to_header: dict[int, str] = {}

    for cell in header_row:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if not name:
            continue
        if name in seen_headers:
            duplicate_headers.add(name)
        else:
            seen_headers.add(name)
            column_order.append(name)
        # Last column with a given name wins for value extraction; the
        # duplicate itself is reported separately so nothing is silently lost.
        col_index_to_header[cell.column] = name

    rows: list[SheetRow] = []
    for row in sheet.iter_rows(min_row=2):
        if all(cell.value is None or cell.value == "" for cell in row):
            continue
        values: dict[str, Any] = {}
        for cell in row:
            header = col_index_to_header.get(cell.column)
            if header is not None:
                values[header] = cell.value
        rows.append(SheetRow(row_number=row[0].row, values=values))

    return ParsedSheet(columns=column_order, duplicate_headers=sorted(duplicate_headers), rows=rows)


def display_value(value: Any) -> str | None:
    """Stringifies a raw cell value for API/report output. `None` and
    empty/whitespace-only strings are both reported as `None` ("blank")."""
    if value is None:
        return None
    if isinstance(value, str):
        return value if value.strip() != "" else None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
